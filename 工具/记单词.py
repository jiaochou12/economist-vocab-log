# -*- coding: utf-8 -*-
"""
记单词 —— 把查到的生词写入单词本.docx（表格：单词 | 词性 | 中文意思 | 查询日期）

用法：
    python 记单词.py lucrative adj. "获利丰厚的；赚大钱的"
    python 记单词.py lucrative adj. "获利丰厚的" --date 2026-09-29

规则（与全局约定一致）：
- 同一单词重复查询不重复记录（按单词小写去重）
- 双写：仓库内 生词表/单词本.docx + 桌面 作业文件/单词本.docx
- 单词本不存在时自动按四列表格格式创建
"""
import argparse
import datetime
import os
import shutil
import sys

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

REPO_WORD_DOC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "生词表", "单词本.docx")
DESKTOP_WORD_DOC = r"D:\HuaweiMoveData\Users\a1505\Desktop\作业文件\单词本.docx"
HEADERS = ["单词", "词性", "中文意思", "查询日期"]


def build_doc(path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc = Document()
    doc.add_paragraph("生词表")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, HEADERS):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
    doc.save(path)
    return doc


def load_doc(path):
    if not os.path.exists(path):
        return build_doc(path)
    return Document(path)


def ensure_table(doc):
    if doc.tables:
        return doc.tables[0]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, HEADERS):
        cell.text = text
    return table


def existing_words(table):
    words = set()
    for row in table.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if cells and cells[0]:
            words.add(cells[0].lower())
    return words


def set_cell_font(cell, text):
    cell.text = text
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.size = Pt(10.5)
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_word(path, word, pos, meaning, date):
    doc = load_doc(path)
    table = ensure_table(doc)
    if word.lower() in existing_words(table):
        return False
    row = table.add_row()
    for cell, text in zip(row.cells, [word, pos, meaning, date]):
        set_cell_font(cell, text)
    doc.save(path)
    return True


def main():
    parser = argparse.ArgumentParser(description="记录生词到单词本.docx")
    parser.add_argument("word", help="单词")
    parser.add_argument("pos", help="词性，如 adj. / n. / v.")
    parser.add_argument("meaning", help="中文意思（可含例句）")
    parser.add_argument("--date", default=datetime.date.today().isoformat(), help="查询日期，默认今天")
    args = parser.parse_args()

    # 仓库副本若不存在而桌面主本存在，先同步一份
    if not os.path.exists(REPO_WORD_DOC) and os.path.exists(DESKTOP_WORD_DOC):
        os.makedirs(os.path.dirname(REPO_WORD_DOC), exist_ok=True)
        shutil.copyfile(DESKTOP_WORD_DOC, REPO_WORD_DOC)

    changed = False
    for path in (os.path.abspath(REPO_WORD_DOC), DESKTOP_WORD_DOC):
        try:
            if add_word(path, args.word, args.pos, args.meaning, args.date):
                changed = True
                print(f"[新增] {path}")
            else:
                print(f"[已存在，跳过] {path}")
        except Exception as exc:  # 桌面不可写等情况不阻断
            print(f"[失败] {path}: {exc}", file=sys.stderr)

    print(f"{'已记录' if changed else '未变更'}: {args.word} | {args.pos} | {args.meaning} | {args.date}")


if __name__ == "__main__":
    main()
